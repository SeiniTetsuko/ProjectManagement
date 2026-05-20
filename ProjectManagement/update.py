#!/usr/bin/env python
# -*- coding: utf-8 -*-

"""
Todo 日报回写应用

作用：
1. 根据 config.todo_project_id 先拉取 Todo 项目下的全量任务，建立 id -> Todo 的索引。
2. 根据 mapping_guid 找到各项目当天日报。
3. 从日报 JSON 中解析 TodoList任务进展：todo_id、承诺完成时间、风险等级、今日进展。
4. 用解析出的 todo_id 匹配 Todo 索引。
5. 匹配成功后，使用该 Todo 的第一个 owner.userGuid 组装 Todo 更新接口 header。
6. 若用户填写的承诺完成时间/风险等级与 Todo 当前值一致，则不重复传该字段，并在 description 中追加提醒。
7. 未填写的字段保持 None，不传入更新 payload。
8. 若用户填写的承诺完成时间/风险等级格式错误，则通过匹配 Todo 的 owner.userGuid 发送飞书个人卡片提醒。
"""

import builtins
import sys
import re
import csv
import json
import time
import traceback
from io import StringIO, BytesIO
from datetime import datetime
from collections import OrderedDict

import requests
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry
from openpyxl import load_workbook

from zenv import get_zdkit_env
from zdbase import ZFile  # 保留平台兼容，不直接依赖

# 如果平台环境里已经内置 send_message_api，则优先复用该函数发送飞书个人卡片。
# 注意：本脚本后续不会定义同名函数，避免覆盖平台能力。
EXTERNAL_SEND_MESSAGE_API = globals().get("send_message_api")

# =============================================================================
# 日志编码兜底
# =============================================================================

try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

# =============================================================================
# print flush patch
# =============================================================================

if not getattr(builtins.print, "_patched_flush", False):
    _original_print = builtins.print

    def print(*args, **kwargs):
        kwargs.setdefault("flush", True)
        _original_print(*args, **kwargs)

    print._patched_flush = True
    builtins.print = print

# =============================================================================
# 固定 API 路由
# =============================================================================

BASE_URL = "https://workspace.cxmt.com"
TODO_BASE_URL = "http://share.cxmt.com"

TODO_LIST_API = "/openapi/todo/todoList"
TODO_UPDATE_API = "/openapi/todo/update"

GET_DOC_API = "/platform/ws/noteInfo/getDocJson"
TOKEN_API = "/api/user/platform/getAccessToken"
DOC_TREE_ROUTE = "/platform/api/main/doc/treeList"
SIGNED_URL_ROUTE = "/platform/api/main/storage/getSignedUrl"

# =============================================================================
# 固定业务参数
# =============================================================================

TODO_STATUS_LIST = ["Open", "Ongoing", "Delay"]
TODO_IS_LEAF = False
TODO_TYPE = "wbs"

RISK_MAP_CN_TO_EN = {
    "高": "high",
    "中": "middle",
    "低": "low",
}

RISK_MAP_EN_TO_CN = {
    "high": "高",
    "middle": "中",
    "low": "低",
}

# 这些内容不应被当作有效进展写回 Todo
DEFAULT_PROGRESS_VALUES = {
    "完成了xxx",
    "开发了xxx",
    "请删除本占位符，并填写今日实际工作进展",
    "请删除本占位符，并填写今日实际完成内容",
    "请在此处填写今日实际工作进展",
    "请在此处填写今日实际完成内容",
}

# =============================================================================
# 全局配置加载
# =============================================================================

zenv_obj = get_zdkit_env()
BASE_URL = zenv_obj.zdkit._http_client.config.get("url")

try:
    with open(config_file.path, "r", encoding="utf-8") as config_fp:
        config = json.load(config_fp)
except Exception as e:
    print(f"❌ 配置文件读取失败: {e}")
    raise

AK = config.get("ak")
SK = config.get("sk")
ORG_GUID = config.get("org_guid")
USER_GUID = config.get("user_guid")

MAPPING_GUID = config.get("mapping_guid")
TARGET_DATE = config.get("target_date") or datetime.now().strftime("%Y-%m-%d")

TODO_PROJECT_ID = config.get("todo_project_id")
if isinstance(TODO_PROJECT_ID, str):
    TODO_PROJECT_ID = [TODO_PROJECT_ID]

# 飞书卡片提醒配置
# 1. 优先复用平台全局 send_message_api（如果存在）
# 2. 如果平台没有全局函数，可在 config 中配置 feishu_message_api_url 作为发送接口
FORMAT_ERROR_CARD_TITLE = config.get(
    "format_error_card_title",
    "WBS任务格式填写错误提醒",
)
FORMAT_ERROR_BUTTON_URL = config.get(
    "format_error_button_url",
    "https://workspace.cxmt.com",
)
FORMAT_ERROR_SENDER_GUID = config.get("format_error_sender_guid") or USER_GUID
FEISHU_MESSAGE_API_URL = (
    config.get("feishu_message_api_url")
    or config.get("message_api_url")
    or ""
)

if not AK or not SK:
    print("❌ 配置错误：缺少 ak / sk")
    raise ValueError("配置错误：缺少 ak / sk")

if not USER_GUID:
    print("❌ 配置错误：缺少 user_guid")
    raise ValueError("配置错误：缺少 user_guid")

if not MAPPING_GUID:
    print("❌ 配置错误：缺少 mapping_guid")
    raise ValueError("配置错误：缺少 mapping_guid")

if not TODO_PROJECT_ID:
    print("❌ 配置错误：缺少 todo_project_id")
    raise ValueError("配置错误：缺少 todo_project_id")

# =============================================================================
# 通用 HTTP 工具
# =============================================================================

session = requests.Session()

retry_strategy = Retry(
    total=5,
    connect=5,
    read=5,
    backoff_factor=2,
    status_forcelist=[429, 500, 502, 503, 504],
    allowed_methods=["HEAD", "GET", "POST", "PUT", "DELETE", "OPTIONS"],
    raise_on_status=False,
)

adapter = HTTPAdapter(
    max_retries=retry_strategy,
    pool_connections=10,
    pool_maxsize=10,
)

session.mount("http://", adapter)
session.mount("https://", adapter)


def request_with_retry(method, url, max_retries=3, **kwargs):
    """
    稳定版 HTTP 请求：
    1. 自动 retry
    2. Connection reset 自动恢复
    3. 避免复用坏连接
    """

    kwargs.setdefault("timeout", 30)
    last_error = None

    for attempt in range(1, max_retries + 1):
        try:
            headers = kwargs.pop("headers", {}) or {}
            headers["Connection"] = "close"
            headers.setdefault("User-Agent", "Mozilla/5.0 PythonRequests")

            if method.lower() == "post":
                response = session.post(url, headers=headers, **kwargs)
            else:
                response = session.get(url, headers=headers, **kwargs)

            return response

        except (
            requests.exceptions.Timeout,
            requests.exceptions.ConnectionError,
            requests.exceptions.ChunkedEncodingError,
        ) as e:
            last_error = e
            wait = min(2 ** attempt, 10)

            print(
                f"⚠️ 请求失败 attempt={attempt}/{max_retries} "
                f"wait={wait}s url={url} error={repr(e)}"
            )

            if attempt < max_retries:
                time.sleep(wait)
            else:
                raise last_error

# =============================================================================
# token / headers
# =============================================================================


def get_access_token():
    print("🔑 获取 AccessToken")

    response = request_with_retry(
        "post",
        BASE_URL + TOKEN_API,
        json={"ak": AK, "sk": SK},
        timeout=30,
    )

    response.raise_for_status()
    response_json = response.json()

    token = (response_json.get("data") or {}).get("accessToken")

    if not token:
        raise Exception(
            f"获取 token 失败: url={BASE_URL + TOKEN_API}, response={response_json}"
        )

    print("✅ AccessToken 获取成功")
    return token


def get_headers(user_guid=None):
    """BASE_URL / Workspace 专用 header。"""
    token = get_access_token()

    return {
        "Access-Token": token,
        "ak": AK,
        "X-User-GUID": user_guid or USER_GUID,
        "Content-Type": "application/json",
    }


def get_todo_headers(user_guid=None):
    """
    TODO_BASE_URL / share 专用 header。

    回写 Todo 时优先使用匹配到的 Todo owner.userGuid。
    拉取 Todo 列表时使用 config.USER_GUID。
    """
    return {
        "x-user-guid": user_guid or USER_GUID,
        "Content-Type": "application/json",
    }

# =============================================================================
# mapping 文件读取
# =============================================================================


def get_signed_url(category_guid):
    print(f"🔗 获取签名 URL: category_guid={category_guid}")

    response = request_with_retry(
        "get",
        BASE_URL + SIGNED_URL_ROUTE,
        headers=get_headers(),
        params={"categoryGuid": category_guid},
        timeout=30,
    )

    response.raise_for_status()
    response_json = response.json()
    signed_url = (response_json.get("data") or {}).get("signedUrl")

    if not signed_url:
        raise Exception(
            f"获取签名 URL 失败: category_guid={category_guid}, response={response_json}"
        )

    return signed_url


def parse_xlsx_mapping(content_bytes):
    workbook = load_workbook(
        filename=BytesIO(content_bytes),
        read_only=True,
        data_only=True,
    )

    sheet = workbook.active
    rows_iter = sheet.iter_rows(values_only=True)

    try:
        headers = next(rows_iter)
    except StopIteration:
        raise ValueError("xlsx 文件为空")

    headers = [str(x).strip() if x is not None else "" for x in headers]
    rows = []

    for excel_row in rows_iter:
        row_dict = {}

        for idx, value in enumerate(excel_row):
            if idx >= len(headers):
                continue

            row_dict[headers[idx]] = str(value).strip() if value is not None else ""

        rows.append({
            "dept": row_dict.get("Dept", ""),
            "project_guid": row_dict.get("project_guid", ""),
            "work_log_folder_guid": row_dict.get("work_log_folder_guid", ""),
            "project_name": row_dict.get("project_name", ""),
        })

    return rows


def parse_csv_mapping(text):
    text = text.lstrip("\ufeff")
    stream = StringIO(text)
    reader = csv.DictReader(stream)

    rows = []

    for row in reader:
        rows.append({
            "dept": row.get("Dept", ""),
            "project_guid": row.get("project_guid", ""),
            "work_log_folder_guid": row.get("work_log_folder_guid", ""),
            "project_name": row.get("project_name", ""),
        })

    return rows


def load_mapping_from_guid(mapping_guid):
    print(f"🚀 开始读取 mapping: {mapping_guid}")

    signed_url = get_signed_url(mapping_guid)

    response = request_with_retry(
        "get",
        signed_url,
        timeout=120,
        stream=True,
        headers={"Connection": "close"},
    )

    response.raise_for_status()
    content = response.content
    signed_url_lower = signed_url.lower()

    if ".xlsx" in signed_url_lower:
        rows = parse_xlsx_mapping(content)
    else:
        rows = parse_csv_mapping(
            content.decode("utf-8-sig", errors="ignore")
        )

    print(f"✅ mapping 读取完成: 有效行数={len(rows)}")
    return rows

# =============================================================================
# Todo API：先按 project_id 拉取全量任务，用于 todo_id 匹配和 header/状态比对
# =============================================================================


def fetch_todo_list():
    print(f"🚀 获取 Todo 列表: projectIds={TODO_PROJECT_ID}")

    request_body = {
        "projectIds": TODO_PROJECT_ID,
        "todoStatus": TODO_STATUS_LIST,
        "isLeaf": TODO_IS_LEAF,
        "type": TODO_TYPE,
    }

    response = request_with_retry(
        "post",
        TODO_BASE_URL + TODO_LIST_API,
        headers=get_todo_headers(),
        json=request_body,
        timeout=60,
    )

    response.raise_for_status()
    response_json = response.json()

    if response_json.get("code") not in (0, 200, "0", "200", None):
        print(
            f"⚠️ Todo API 返回 code 异常: "
            f"code={response_json.get('code')}, msg={response_json.get('msg')}"
        )

    todo_list = response_json.get("data") or []
    print(f"✅ 获取 Todo 数量: {len(todo_list)}")

    return todo_list


def build_todo_index(todo_list):
    """构建 id -> todo 的索引。"""
    todo_index = {}
    duplicate_ids = set()

    for todo in todo_list:
        todo_id = todo.get("id")
        if todo_id in (None, ""):
            continue

        todo_id_int = safe_int(todo_id)
        if todo_id_int is None:
            continue

        if todo_id_int in todo_index:
            duplicate_ids.add(todo_id_int)

        todo_index[todo_id_int] = todo

    if duplicate_ids:
        print(f"⚠️ Todo id 存在重复，已以后出现者覆盖: {sorted(duplicate_ids)}")

    print(f"✅ Todo 索引构建完成: {len(todo_index)} 条")
    return todo_index


def safe_int(value):
    try:
        return int(value)
    except Exception:
        return None


def get_first_owner_user_guid(todo):
    owners = todo.get("owners", []) or []
    for owner in owners:
        user_guid = owner.get("userGuid")
        if user_guid:
            return user_guid
    return None


def get_first_owner_display_name(todo):
    owners = todo.get("owners", []) or []
    for owner in owners:
        return owner.get("displayName") or owner.get("name") or owner.get("loginName") or ""
    return ""


def get_todo_code_for_display(todo):
    """用于卡片展示的 Todo code；接口缺失 code 时用 WBS#{id} 兜底。"""
    code = str(todo.get("code") or "").strip()
    if code:
        return code

    todo_id = todo.get("id")
    if todo_id not in (None, ""):
        return f"WBS#{todo_id}"

    return ""


def get_todo_related_project_name(todo, fallback_project_name=""):
    related_project = todo.get("relatedProject") or {}
    return (
        related_project.get("name")
        or fallback_project_name
        or "未命名项目"
    )


def normalize_date_for_compare(date_value):
    text = str(date_value or "").strip()
    if not text:
        return ""

    m = re.match(r"^(\d{4}-\d{2}-\d{2})", text)
    if m:
        return m.group(1)

    m = re.match(r"^(\d{4})/(\d{1,2})/(\d{1,2})", text)
    if m:
        y, mo, d = m.groups()
        return f"{int(y):04d}-{int(mo):02d}-{int(d):02d}"

    m = re.match(r"^(\d{4})(\d{2})(\d{2})$", text)
    if m:
        y, mo, d = m.groups()
        return f"{y}-{mo}-{d}"

    return text

# =============================================================================
# 日报查找
# =============================================================================


def get_tree_node_title(node):
    for key in ("dataTitle", "title", "name", "fileName", "filename"):
        value = node.get(key)
        if value:
            return str(value).strip()
    return ""


def get_tree_node_guid(node):
    for key in ("categoryGuid", "dataGuid", "guid", "fileGuid", "id"):
        value = node.get(key)
        if value:
            return str(value)
    return ""


def get_date_title_variants(date_str):
    dt = datetime.strptime(date_str, "%Y-%m-%d")

    return [
        dt.strftime("%Y-%m-%d"),
        dt.strftime("%Y/%m/%d"),
        dt.strftime("%Y%m%d"),
        dt.strftime("%Y.%m.%d"),
        dt.strftime("%Y年%m月%d日"),
        f"{dt.year}年{dt.month}月{dt.day}日",
    ]


def is_daily_note_title(title, target_date):
    title = title or ""

    has_date = any(x in title for x in get_date_title_variants(target_date))
    if not has_date:
        return False

    return "日报" in title


def score_daily_note_title(title, target_date, project_name=""):
    score = 0
    title = title or ""

    if is_daily_note_title(title, target_date):
        score += 100

    if project_name and project_name in title:
        score += 30

    if "副本" in title:
        score -= 20

    if "测试" in title:
        score -= 10

    if title.endswith("日报"):
        score += 5

    return score


def list_folder_nodes(project_guid, folder_guid):
    response = request_with_retry(
        "post",
        BASE_URL + DOC_TREE_ROUTE,
        headers=get_headers(),
        json={
            "projectGuid": project_guid,
            "parentGuid": folder_guid,
        },
        timeout=60,
    )

    response.raise_for_status()
    response_json = response.json()

    return response_json.get("data") or []


def find_today_daily_note(project_guid, folder_guid, target_date, project_name=""):
    print(
        f"🔎 查找日报: project_guid={project_guid}, "
        f"folder_guid={folder_guid}, date={target_date}, project_name={project_name}"
    )

    nodes = list_folder_nodes(project_guid, folder_guid)
    candidates = []

    for node in nodes:
        title = get_tree_node_title(node)
        guid = get_tree_node_guid(node)

        if not title or not guid:
            continue

        if not is_daily_note_title(title, target_date):
            continue

        candidates.append({
            "note_guid": guid,
            "note_title": title,
            "score": score_daily_note_title(title, target_date, project_name),
            "node": node,
        })

    if not candidates:
        return None

    candidates.sort(key=lambda x: x["score"], reverse=True)
    best = candidates[0]

    print(
        f"✅ 命中日报: {best['note_title']} | "
        f"note_guid={best['note_guid']} | score={best['score']}"
    )

    if len(candidates) > 1:
        print("⚠️ 存在多个候选日报，已选择 score 最高的一个：")
        for item in candidates:
            print(f"    - score={item['score']} | {item['note_title']} | {item['note_guid']}")

    return best

# =============================================================================
# note json
# =============================================================================


def get_doc_json(doc_id):
    print(f"📄 获取日报内容: doc_id={doc_id}")

    response = request_with_retry(
        "get",
        BASE_URL + GET_DOC_API,
        headers=get_headers(),
        params={"docId": doc_id},
        timeout=60,
    )

    response.raise_for_status()
    return response.json()


def extract_block_text(block):
    """从单个 blockContainer 提取纯文本。"""
    texts = []
    _collect_text_from_node(block, texts)
    return "".join(texts).strip()


def _collect_text_from_node(node, texts):
    """递归提取 text 节点内容，忽略 mention 等非文字节点。"""
    if isinstance(node, list):
        for item in node:
            _collect_text_from_node(item, texts)
        return

    if not isinstance(node, dict):
        return

    if node.get("type") == "text":
        text = node.get("text", "")
        if text:
            texts.append(text)

    for child in node.get("content", []) or []:
        _collect_text_from_node(child, texts)


def _collect_block_containers(node, result):
    """递归收集所有 blockContainer 节点，保持文档顺序。"""
    if isinstance(node, list):
        for item in node:
            _collect_block_containers(item, result)
        return

    if not isinstance(node, dict):
        return

    if node.get("type") == "blockContainer":
        result.append(node)
        return

    for child in node.get("content", []) or []:
        _collect_block_containers(child, result)


def _get_block_primary_type(block_container):
    """获取 blockContainer 下第一个内容节点的 type。"""
    for child in block_container.get("content", []) or []:
        child_type = child.get("type", "")
        if child_type:
            return child_type
    return ""


def _get_heading_level(block_container):
    """获取 heading 的 level 属性。"""
    for child in block_container.get("content", []) or []:
        if child.get("type") == "heading":
            attrs = child.get("attrs", {}) or {}
            level = attrs.get("level", "")
            return str(level)
    return ""

# =============================================================================
# progress 判定
# =============================================================================


def normalize_progress_line(line):
    line = str(line or "").strip()
    line = re.sub(r"^[\-\•\*]+[\s]*", "", line).strip()
    line = re.sub(r"^\d+[.\)]\s*", "", line).strip()
    return line


def is_meaningful_progress(progress):
    if not progress:
        return False

    lines = []

    for line in progress.splitlines():
        line = normalize_progress_line(line)

        if not line:
            continue

        if line in DEFAULT_PROGRESS_VALUES:
            continue

        lines.append(line)

    return len(lines) > 0

# =============================================================================
# task parse from doc json
# =============================================================================


def parse_tasks_from_doc(doc_json):
    """
    从笔记 JSON 结构解析 Todo 任务。

    兼容刚才写入日报应用生成的格式：
    - 🚀 【项目名】&【任务标题】&【50001】
    - <任务截止时间>: 2026-05-20  只展示，不回写
    - <承诺完成时间：YYYY-MM-DD>: 2026-05-21
    - <风险等级：高/中/低>: 中
    - <今日进展描述>:
      - 实际进展
    """

    content = (doc_json.get("data") or {}).get("content", [])

    all_blocks = []
    _collect_block_containers(content, all_blocks)

    tasks = []
    skip_default_count = 0

    print("🚀 开始解析日报内容（JSON 结构）")

    current_task = None

    for block in all_blocks:
        block_text = extract_block_text(block)
        primary_type = _get_block_primary_type(block)

        # h2/h3 是任务边界
        if primary_type == "heading":
            level = _get_heading_level(block)

            if level in ("2", "3"):
                if current_task:
                    skip_default_count = _finalize_task(
                        current_task, tasks, skip_default_count
                    )
                    current_task = None

                continue

        # 🚀 开头 = 新的 Todo 标题行
        if "🚀" in block_text:
            if current_task:
                skip_default_count = _finalize_task(
                    current_task, tasks, skip_default_count
                )

            # 取标题中的最后一个纯数字【xxx】作为 todo_id。
            # 这样即使标题/项目名中误含数字，也优先取末尾 id。
            todo_matches = re.findall(r"【(\d+)】", block_text)
            todo_id = safe_int(todo_matches[-1]) if todo_matches else None

            current_task = {
                "todo_id": todo_id,
                "title_line": block_text,

                # 承诺完成时间
                "promise_raw": None,
                "promise_date": None,
                "promise_invalid": False,

                # 风险等级
                "risk_raw": None,
                "risk_level": None,
                "risk_invalid": False,

                # 今日进展
                "progress_lines": [],
                "in_progress_section": False,
            }

            continue

        # 还没遇到 🚀，跳过
        if not current_task:
            continue

        # 任务截止时间仅展示，不回写
        due_match = re.match(r"<任务截止时间[^>]*>:[ \t]*(.*)", block_text)
        if due_match:
            current_task["in_progress_section"] = False
            continue

        # 承诺完成时间：兼容 <承诺完成时间：YYYY-MM-DD>: 2026-05-21
        promise_match = re.match(r"<承诺完成时间[^>]*>:[ \t]*(.*)", block_text)
        if promise_match:
            user_value = promise_match.group(1).strip()
            current_task["promise_raw"] = user_value

            date_match = re.fullmatch(r"\d{4}-\d{2}-\d{2}", user_value)
            if date_match:
                current_task["promise_date"] = user_value
            elif user_value:
                current_task["promise_invalid"] = True

            current_task["in_progress_section"] = False
            continue

        # 风险等级：兼容 <风险等级：高/中/低>: 中
        risk_match = re.match(r"<风险等级[^>]*>:[ \t]*(.*)", block_text)
        if risk_match:
            user_value = risk_match.group(1).strip()
            current_task["risk_raw"] = user_value

            if user_value in RISK_MAP_CN_TO_EN:
                current_task["risk_level"] = RISK_MAP_CN_TO_EN[user_value]
            elif user_value:
                current_task["risk_invalid"] = True

            current_task["in_progress_section"] = False
            continue

        # 今日进展描述：兼容 <今日进展描述> 或 <今日进展描述：2026-05-20>
        progress_match = re.match(r"<今日进展描述[^>]*>[:：]?[ \t]*(.*)", block_text)
        if progress_match:
            current_task["in_progress_section"] = True
            inline_text = progress_match.group(1).strip()

            if inline_text:
                line = normalize_progress_line(inline_text)
                if line:
                    current_task["progress_lines"].append(line)

            continue

        # 收集进展内容（bulletListItem / paragraph）
        if current_task["in_progress_section"]:
            line = normalize_progress_line(block_text)
            if line:
                current_task["progress_lines"].append(line)

    # 保存最后一个任务
    if current_task:
        skip_default_count = _finalize_task(current_task, tasks, skip_default_count)

    print(
        f"✅ 解析完成: 有效任务={len(tasks)}, "
        f"跳过默认模板={skip_default_count}"
    )

    return tasks, skip_default_count


def _finalize_task(current_task, tasks, skip_default_count):
    """收尾一个 task，只负责从日报内容生成待匹配的解析结果。"""

    if not current_task.get("todo_id"):
        print(f"⚠️ 跳过无法解析 todo_id 的任务行: {current_task.get('title_line', '')}")
        return skip_default_count

    raw_progress = "\n".join(current_task["progress_lines"]).strip()
    has_meaningful_progress = is_meaningful_progress(raw_progress)

    warnings = []

    if current_task["promise_invalid"]:
        warnings.append("- 承诺完成时间填写格式有误，正确格式：YYYY-MM-DD")

    if current_task["risk_invalid"]:
        warnings.append("- 风险等级填写格式有误，正确格式：高/中/低")

    has_promise_input = bool(current_task["promise_raw"])
    has_risk_input = bool(current_task["risk_raw"])

    has_any_update = (
        has_meaningful_progress
        or has_promise_input
        or has_risk_input
        or bool(warnings)
    )

    if not has_any_update:
        print(f"ℹ️ 跳过默认模板内容 todo={current_task['todo_id']}")
        return skip_default_count + 1

    progress_parts = []

    if warnings:
        progress_parts.append(f"【格式检查提醒 {TARGET_DATE}】")
        progress_parts.extend(warnings)

    if has_meaningful_progress:
        if progress_parts:
            progress_parts.append("")
        progress_parts.append(raw_progress)

    progress = "\n".join(progress_parts).strip()

    tasks.append({
        "id": current_task["todo_id"],

        # 用户没写或格式错误时为 None
        "promiseDate": current_task["promise_date"],
        "promiseRaw": current_task["promise_raw"],
        "promiseInvalid": current_task["promise_invalid"],

        # 用户没写或格式错误时为 None
        "riskLevel": current_task["risk_level"],
        "riskRaw": current_task["risk_raw"],
        "riskInvalid": current_task["risk_invalid"],

        "progress": progress,
    })

    return skip_default_count

# =============================================================================
# 与 Todo 当前状态比对，并构造更新 payload
# =============================================================================


def enrich_task_with_todo_context(parsed_task, todo_index, note_project_name=""):
    """
    用解析出的 todo_id 匹配项目 Todo 列表。

    处理规则：
    - 匹配不到：返回 None，由主流程跳过更新。
    - 匹配到：取第一个 owner.userGuid 作为 Todo 更新 header 的 x-user-guid。
    - promiseDate/riskLevel 若用户没写：保持 None，不传入接口。
    - promiseDate/riskLevel 若用户写了且与 Todo 当前值一致：不传入接口，description 追加一致提醒。
    - promiseDate/riskLevel 若用户写了且与 Todo 当前值不同：传入接口更新。
    """

    todo_id = parsed_task["id"]
    source_todo = todo_index.get(todo_id)

    if not source_todo:
        print(
            f"⚠️ 日报中解析到 todo_id={todo_id}，"
            f"但未在 config.todo_project_id 拉取的 Todo 列表中找到，跳过更新"
        )
        return None

    owner_user_guid = get_first_owner_user_guid(source_todo)
    owner_display_name = get_first_owner_display_name(source_todo)
    source_project_name = get_todo_related_project_name(
        source_todo,
        fallback_project_name=note_project_name,
    )
    todo_code = get_todo_code_for_display(source_todo)

    if not owner_user_guid:
        print(
            f"⚠️ Todo 匹配成功但缺少 owners.userGuid，跳过更新: "
            f"todo_id={todo_id}, title={source_todo.get('title', '')}"
        )
        return None

    current_promise = normalize_date_for_compare(source_todo.get("promiseDate"))
    current_risk = str(source_todo.get("riskLevel") or "").strip()

    update_promise = parsed_task.get("promiseDate")
    update_risk = parsed_task.get("riskLevel")

    unchanged_messages = []

    # 承诺完成时间：用户填写了合法日期，才进行比对
    if update_promise is not None:
        if normalize_date_for_compare(update_promise) == current_promise:
            unchanged_messages.append("- 承诺完成时间和上次更新一致")
            update_promise = None

    # 风险等级：用户填写了合法枚举，才进行比对
    if update_risk is not None:
        if update_risk == current_risk:
            unchanged_messages.append("- 风险等级和上次更新一致")
            update_risk = None

    format_errors = []

    if parsed_task.get("promiseInvalid"):
        format_errors.append({
            "field": "承诺完成时间",
            "raw": parsed_task.get("promiseRaw") or "空",
            "expected": "YYYY-MM-DD",
        })

    if parsed_task.get("riskInvalid"):
        format_errors.append({
            "field": "风险等级",
            "raw": parsed_task.get("riskRaw") or "空",
            "expected": "高/中/低",
        })

    description_parts = []

    progress = parsed_task.get("progress", "").strip()
    if progress:
        description_parts.append(progress)

    if unchanged_messages:
        if description_parts:
            description_parts.append("")
        description_parts.append(f"【字段一致提醒 {TARGET_DATE}】")
        description_parts.extend(unchanged_messages)

    description = "\n".join(description_parts).strip()

    # 允许只写 promise/risk 不写 progress。
    # 如果用户只写了与现值一致的字段，也会产生 description 一致提醒。
    if not description and update_promise is None and update_risk is None:
        print(
            f"ℹ️ Todo 无有效变化，跳过更新: "
            f"todo_id={todo_id}, title={source_todo.get('title', '')}"
        )
        return None

    return {
        "id": todo_id,
        "type": TODO_TYPE,
        "description": description,
        "promiseDate": update_promise,
        "riskLevel": update_risk,
        "owner_user_guid": owner_user_guid,
        "owner_display_name": owner_display_name,
        "source_todo_title": source_todo.get("title", ""),
        "source_project_name": source_project_name,
        "todo_code": todo_code,
        "format_errors": format_errors,
        "currentPromiseDate": current_promise,
        "currentRiskLevel": current_risk,
    }

# =============================================================================
# 飞书格式错误提醒卡片
# =============================================================================


def build_format_error_card_content(task):
    """构造格式错误卡片正文。"""
    project_name = task.get("source_project_name", "未命名项目")
    todo_name = task.get("source_todo_title", "未命名任务")
    todo_code = task.get("todo_code", "")

    task_line = f"【{project_name}】【{todo_name}】【{todo_code}】"

    lines = [
        f"**您的任务：**{task_line} 格式填写有误",
        "",
        "**错误详情：**",
    ]

    for error in task.get("format_errors", []) or []:
        field = error.get("field", "")
        raw_value = error.get("raw", "")
        expected = error.get("expected", "")
        lines.append(f"- **<{field}>** 填写为：`{raw_value}`")
        lines.append(f"  应填写格式：`{expected}`")

    lines.extend([
        "",
        "请前往项目中心及时更新。",
    ])

    return "\n".join(lines)


def build_format_error_plain_text(task):
    """构造普通文本消息，作为 interactive_content 外的兜底内容。"""
    project_name = task.get("source_project_name", "未命名项目")
    todo_name = task.get("source_todo_title", "未命名任务")
    todo_code = task.get("todo_code", "")

    error_parts = []
    for error in task.get("format_errors", []) or []:
        field = error.get("field", "")
        raw_value = error.get("raw", "")
        expected = error.get("expected", "")
        error_parts.append(f"<{field}>填写为：{raw_value}，应填写格式为：{expected}")

    error_text = "；".join(error_parts)

    return (
        f"您的任务：【{project_name}】【{todo_name}】【{todo_code}】格式填写有误。"
        f"{error_text}。请前往项目中心及时更新：{FORMAT_ERROR_BUTTON_URL}"
    )


def build_format_error_feishu_card(task):
    """
    构造飞书卡片。
    参考用户提供的 schema 2.0 卡片结构，但这里只保留一个“前往项目中心”按钮。
    """
    card_content = build_format_error_card_content(task)

    return {
        "schema": "2.0",
        "header": {
            "padding": "12px 8px 12px 8px",
            "template": "red",
            "title": {
                "content": FORMAT_ERROR_CARD_TITLE,
                "tag": "plain_text",
            },
        },
        "body": {
            "vertical_spacing": "12px",
            "elements": [
                {
                    "tag": "markdown",
                    "content": card_content,
                    "margin": "0px",
                    "text_size": "normal",
                },
                {
                    "tag": "column_set",
                    "flex_mode": "stretch",
                    "horizontal_spacing": "8px",
                    "margin": "0px",
                    "columns": [
                        {
                            "tag": "column",
                            "width": "weighted",
                            "weight": 1,
                            "elements": [
                                {
                                    "tag": "button",
                                    "type": "primary_filled",
                                    "width": "fill",
                                    "margin": "4px 0px 4px 0px",
                                    "text": {
                                        "tag": "plain_text",
                                        "content": "前往项目中心",
                                    },
                                    "behaviors": [
                                        {
                                            "type": "open_url",
                                            "default_url": FORMAT_ERROR_BUTTON_URL,
                                        }
                                    ],
                                }
                            ],
                        }
                    ],
                },
            ],
        },
    }


def send_format_error_card(task):
    """
    给 Todo 第一个 owner.userGuid 发送格式错误提醒卡片。

    发送策略：
    1. 如果平台环境存在全局 EXTERNAL_SEND_MESSAGE_API，则优先调用。
    2. 否则如果配置了 feishu_message_api_url/message_api_url，则直接 POST。
    3. 两者都没有时，只打印日志，不中断 Todo 回写。
    """
    format_errors = task.get("format_errors", []) or []
    if not format_errors:
        return False

    receiver_guid = task.get("owner_user_guid")
    if not receiver_guid:
        print(f"⚠️ 格式错误卡片未发送：缺少 receiver owner_user_guid, todo={task.get('id')}")
        return False

    card = build_format_error_feishu_card(task)
    plain_text = build_format_error_plain_text(task)

    print(
        f"📩 准备发送格式错误提醒卡片: "
        f"todo={task.get('id')}, receiver={receiver_guid}"
    )

    # 1. 优先复用平台已有函数
    if callable(EXTERNAL_SEND_MESSAGE_API):
        try:
            response = EXTERNAL_SEND_MESSAGE_API(
                receiver_guids=[receiver_guid],
                title=FORMAT_ERROR_CARD_TITLE,
                content=plain_text,
                sender_guid=FORMAT_ERROR_SENDER_GUID,
                interactive_content=card,
            )

            if hasattr(response, "status_code"):
                ok = response.status_code == 200
                try:
                    response_json = response.json()
                except Exception:
                    response_json = {}

                if ok and (response_json.get("data") or response_json.get("code") in (0, 200, "0", "200", None)):
                    print(f"✅ 格式错误提醒卡片发送成功: todo={task.get('id')}")
                    return True

                print(f"❌ 格式错误提醒卡片发送失败: status={response.status_code}, text={getattr(response, 'text', '')}")
                return False

            print(f"✅ 格式错误提醒卡片发送完成: todo={task.get('id')}, response={response}")
            return True

        except Exception as e:
            print(f"❌ 调用平台 send_message_api 发送卡片异常: todo={task.get('id')}, error={e}")
            print(traceback.format_exc())
            return False

    # 2. 配置了直接发送接口时使用 HTTP POST
    if FEISHU_MESSAGE_API_URL:
        try:
            payload = {
                "receiver_guids": [receiver_guid],
                "title": FORMAT_ERROR_CARD_TITLE,
                "content": plain_text,
                "sender_guid": FORMAT_ERROR_SENDER_GUID,
                "interactive_content": card,
            }

            response = request_with_retry(
                "post",
                FEISHU_MESSAGE_API_URL,
                headers=get_headers(user_guid=FORMAT_ERROR_SENDER_GUID),
                json=payload,
                timeout=30,
            )

            if response.status_code != 200:
                print(
                    f"❌ 格式错误提醒卡片发送失败: "
                    f"status={response.status_code}, text={response.text[:500]}"
                )
                return False

            response_json = response.json()
            if response_json.get("data") or response_json.get("code") in (0, 200, "0", "200", None):
                print(f"✅ 格式错误提醒卡片发送成功: todo={task.get('id')}")
                return True

            print(f"❌ 格式错误提醒卡片发送失败: response={response_json}")
            return False

        except Exception as e:
            print(f"❌ 通过 feishu_message_api_url 发送卡片异常: todo={task.get('id')}, error={e}")
            print(traceback.format_exc())
            return False

    # 3. 没有发送能力时跳过，不影响回写
    print(
        "⚠️ 未检测到平台 send_message_api，且未配置 feishu_message_api_url/message_api_url，"
        f"跳过格式错误提醒卡片发送: todo={task.get('id')}"
    )
    return False


# =============================================================================
# update todo
# =============================================================================


def update_todo(task):
    payload = {
        "type": TODO_TYPE,
        "id": task["id"],
        "description": task["description"],
    }

    # 没写 / 与当前值一致 / 格式错误 -> None，不传入接口
    if task.get("promiseDate") is not None:
        payload["promiseDate"] = task["promiseDate"]

    if task.get("riskLevel") is not None:
        payload["riskLevel"] = task["riskLevel"]

    owner_user_guid = task.get("owner_user_guid")
    owner_display_name = task.get("owner_display_name", "")

    print(
        f"\n🚀 更新 Todo: id={task['id']}, "
        f"owner={owner_display_name}, owner_user_guid={owner_user_guid}"
    )
    print(json.dumps(payload, ensure_ascii=False, indent=2))

    response = request_with_retry(
        "post",
        TODO_BASE_URL + TODO_UPDATE_API,
        headers=get_todo_headers(user_guid=owner_user_guid),
        json=payload,
        timeout=30,
    )

    response.raise_for_status()
    result = response.json()

    if result.get("code") not in (0, 200, "0", "200", None):
        raise Exception(f"Todo 更新接口返回异常: response={result}")

    print(f"✅ 更新成功: todo={task['id']}")
    return result

# =============================================================================
# 主流程
# =============================================================================

print("\n==========================================")
print("🚀 Todo 日报回写应用启动")
print("==========================================\n")

print(f"BASE_URL: {BASE_URL}")
print(f"TODO_BASE_URL: {TODO_BASE_URL}")
print(f"TARGET_DATE: {TARGET_DATE}")
print(f"USER_GUID: {USER_GUID}")
print(f"MAPPING_GUID: {MAPPING_GUID}")
print(f"TODO_PROJECT_ID: {TODO_PROJECT_ID}")

mapping_rows = load_mapping_from_guid(MAPPING_GUID)

todo_list = fetch_todo_list()
if not todo_list:
    print("ℹ️ 当前未获取到 Todo 数据，流程结束")
    sys.exit(0)

todo_index = build_todo_index(todo_list)

visited_notes = set()

total_projects = 0
hit_notes = 0
total_parsed_tasks = 0
total_update_tasks = 0
success_count = 0
fail_count = 0
skip_default_total = 0
skip_no_match_count = 0
skip_no_change_count = 0
card_success_count = 0
card_fail_count = 0
card_skip_count = 0

for row in mapping_rows:
    project_guid = row.get("project_guid")
    folder_guid = row.get("work_log_folder_guid")
    project_name = row.get("project_name", "")

    if not project_guid:
        continue

    if not folder_guid:
        continue

    total_projects += 1

    print(f"\n🚀 处理项目: {project_name}")

    daily_note = find_today_daily_note(
        project_guid=project_guid,
        folder_guid=folder_guid,
        target_date=TARGET_DATE,
        project_name=project_name,
    )

    if not daily_note:
        print("ℹ️ 未找到今日日报")
        continue

    note_guid = daily_note["note_guid"]

    if note_guid in visited_notes:
        print(f"ℹ️ 日报已处理过，跳过重复 note: {note_guid}")
        continue

    visited_notes.add(note_guid)
    hit_notes += 1

    print(f"✅ 命中日报: {daily_note['note_title']}")

    doc_json = get_doc_json(note_guid)

    parsed_tasks, skip_default_count = parse_tasks_from_doc(doc_json)
    skip_default_total += skip_default_count

    if not parsed_tasks:
        print("ℹ️ 当前日报无有效 Todo 更新")
        continue

    total_parsed_tasks += len(parsed_tasks)

    print("\n======================")
    print("日报解析结果")
    print("======================")
    print(json.dumps(parsed_tasks, ensure_ascii=False, indent=2))

    update_tasks = []

    for parsed_task in parsed_tasks:
        enriched_task = enrich_task_with_todo_context(
            parsed_task,
            todo_index,
            note_project_name=project_name,
        )

        if not enriched_task:
            # 粗略统计：匹配不到和无变化都在 enrich 内部打印，外部按 todo_index 区分一次
            if parsed_task.get("id") not in todo_index:
                skip_no_match_count += 1
            else:
                skip_no_change_count += 1
            continue

        update_tasks.append(enriched_task)

    if not update_tasks:
        print("ℹ️ 当前日报没有需要写回 Todo 的有效更新")
        continue

    total_update_tasks += len(update_tasks)

    print("\n======================")
    print("待更新 Todo 结果")
    print("======================")
    print(json.dumps(update_tasks, ensure_ascii=False, indent=2))

    for task in update_tasks:
        # 如果承诺完成时间/风险等级格式错误，先给该 Todo owner 发送飞书卡片提醒。
        # 卡片发送失败不阻断 Todo description 写回。
        if task.get("format_errors"):
            sent = send_format_error_card(task)
            if sent:
                card_success_count += 1
            else:
                card_fail_count += 1
        else:
            card_skip_count += 1

        try:
            update_todo(task)
            success_count += 1
        except Exception as e:
            fail_count += 1
            print(f"❌ 更新失败 todo={task['id']} error={e}")
            print(traceback.format_exc())

print("\n==========================================")
print("🏁 Todo 日报回写应用完成")
print("==========================================\n")

print("================ 执行统计 ================")
print(f"处理项目数: {total_projects}")
print(f"命中日报数: {hit_notes}")
print(f"解析任务数: {total_parsed_tasks}")
print(f"待更新任务数: {total_update_tasks}")
print(f"更新成功: {success_count}")
print(f"更新失败: {fail_count}")
print(f"跳过(默认模板): {skip_default_total}")
print(f"跳过(未匹配 Todo): {skip_no_match_count}")
print(f"跳过(无有效变化/缺少owner): {skip_no_change_count}")
print(f"格式错误卡片发送成功: {card_success_count}")
print(f"格式错误卡片发送失败/跳过: {card_fail_count}")
print(f"无格式错误无需发卡片: {card_skip_count}")
print("==========================================\n")

if fail_count > 0:
    raise Exception(f"存在 {fail_count} 个 Todo 更新失败，请查看日志")
